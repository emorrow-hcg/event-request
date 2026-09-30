"""Generate SOP Word documents for PyCharm and VS Code setup."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy


# ── Colour palette ────────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1F, 0x35, 0x64)   # title / headings
MID_BLUE   = RGBColor(0x2E, 0x74, 0xB5)   # section headers
ACCENT     = RGBColor(0x00, 0x70, 0xC0)   # step numbers / links
TIP_BG     = RGBColor(0xE2, 0xEF, 0xDA)   # tip box background (light green)
WARN_BG    = RGBColor(0xFF, 0xE5, 0x99)   # warning box (light yellow)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
CODE_GREY  = RGBColor(0x26, 0x26, 0x26)
CODE_BG    = RGBColor(0xF0, 0xF0, 0xF0)


def _hex(colour: RGBColor) -> str:
    return f'{colour[0]:02X}{colour[1]:02X}{colour[2]:02X}'


def set_cell_bg(cell, colour: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), _hex(colour))
    tcPr.append(shd)


def add_heading(doc, text, level=1, colour=DARK_BLUE):
    p = doc.add_paragraph()
    p.style = f'Heading {level}'
    run = p.add_run(text)
    run.font.color.rgb = colour
    if level == 1:
        run.font.size = Pt(20)
        run.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = MID_BLUE
    elif level == 3:
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = ACCENT
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style = 'Normal'
    for run in p.runs:
        run.font.size = Pt(11)
    return p


def add_step(doc, number, title, details=None):
    """Add a numbered step with bold title and optional detail lines."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(6)
    num_run = p.add_run(f"Step {number}: ")
    num_run.bold = True
    num_run.font.color.rgb = ACCENT
    num_run.font.size = Pt(11)
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(11)
    if details:
        for line in details:
            dp = doc.add_paragraph(line)
            dp.paragraph_format.left_indent = Inches(0.5)
            dp.paragraph_format.space_before = Pt(2)
            dp.paragraph_format.space_after = Pt(2)
            for r in dp.runs:
                r.font.size = Pt(10.5)


def add_code(doc, code_text):
    """Add a shaded code block."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    # shade the paragraph
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F0F0')
    pPr.append(shd)
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9.5)
    run.font.color.rgb = CODE_GREY
    return p


def add_tip(doc, text, kind='TIP'):
    """Add a tip or warning callout box using a 1-column table."""
    colour = TIP_BG if kind == 'TIP' else WARN_BG
    label_colour = RGBColor(0x37, 0x5E, 0x23) if kind == 'TIP' else RGBColor(0x7F, 0x60, 0x00)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, colour)
    p = cell.paragraphs[0]
    label = p.add_run(f"  {kind}:  ")
    label.bold = True
    label.font.size = Pt(10)
    label.font.color.rgb = label_colour
    body = p.add_run(text)
    body.font.size = Pt(10)
    doc.add_paragraph()  # spacer


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3 + level * 0.2)
    for r in p.runs:
        r.font.size = Pt(10.5)
    return p


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2E74B5')
    pBdr.append(bottom)
    pPr.append(pBdr)


def title_page(doc, title, subtitle, ide_name):
    # Big title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(28)
    run.bold = True
    run.font.color.rgb = DARK_BLUE

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(14)
    r2.font.color.rgb = MID_BLUE

    doc.add_paragraph()

    # IDE badge
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, MID_BLUE)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(f"  {ide_name}  ")
    cr.font.size = Pt(16)
    cr.bold = True
    cr.font.color.rgb = WHITE

    doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        "ATS Resume Optimizer  |  Setup & Usage Guide\n"
        "Repository: github.com/emorrow-hcg/event-request\n"
        "Python 3.13+  |  Version 1.0"
    )
    mr.font.size = Pt(10)
    mr.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════════════
# PYCHARM SOP
# ═══════════════════════════════════════════════════════════════════════════════

def build_pycharm_sop():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    title_page(
        doc,
        "ATS Resume Optimizer",
        "Step-by-Step Setup & Usage Guide",
        "PyCharm Edition"
    )

    # ── SECTION 1: Overview ──────────────────────────────────────────────────
    add_heading(doc, "1. Overview", level=1)
    add_body(doc,
        "This guide walks you through every step needed to get the ATS Resume Optimizer "
        "running on your computer using PyCharm — one of the most popular Python IDEs. "
        "No prior experience is required. Follow each step in order and you will have the "
        "application running in under 20 minutes."
    )
    add_tip(doc,
        "ATS stands for Applicant Tracking System. This tool reads a job description, "
        "scores your resume against it, and uses AI to rewrite your resume so it scores higher.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 2: Prerequisites ─────────────────────────────────────────────
    add_heading(doc, "2. What You Need Before You Start", level=1)
    add_body(doc, "You will need to install the following software before opening PyCharm:")

    prereqs = [
        ("Python 3.13 (latest)",
         "https://www.python.org/downloads/",
         "Download and run the installer. On Windows, tick 'Add Python to PATH' before clicking Install."),
        ("Git",
         "https://git-scm.com/downloads",
         "Needed to download (clone) the project code from GitHub."),
        ("PyCharm Community Edition (free)",
         "https://www.jetbrains.com/pycharm/download/",
         "Choose the Community edition — it is free and has everything you need."),
        ("An Anthropic API Key",
         "https://console.anthropic.com",
         "Sign in, go to API Keys, click Create Key, and copy the key somewhere safe."),
    ]

    for i, (name, url, note) in enumerate(prereqs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        nr = p.add_run(f"  {i}.  ")
        nr.bold = True
        nr.font.color.rgb = ACCENT
        nr.font.size = Pt(11)
        br = p.add_run(name)
        br.bold = True
        br.font.size = Pt(11)
        urlp = doc.add_paragraph()
        urlp.paragraph_format.left_indent = Inches(0.6)
        urlp.paragraph_format.space_before = Pt(0)
        urlp.paragraph_format.space_after = Pt(2)
        ur = urlp.add_run(f"Download: {url}")
        ur.font.size = Pt(10)
        ur.font.color.rgb = ACCENT
        np_ = doc.add_paragraph()
        np_.paragraph_format.left_indent = Inches(0.6)
        np_.paragraph_format.space_before = Pt(0)
        np_.paragraph_format.space_after = Pt(8)
        nr2 = np_.add_run(note)
        nr2.font.size = Pt(10)

    add_tip(doc,
        "To check if Python is installed correctly, open a terminal (Command Prompt on Windows, "
        "Terminal on Mac/Linux) and type:  python --version  or  python3 --version  "
        "You should see 'Python 3.13.x'.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 3: Getting the Code ──────────────────────────────────────────
    add_heading(doc, "3. Getting the Project Code", level=1)
    add_body(doc,
        "You need to copy (clone) the project from GitHub to your computer. "
        "You can do this directly inside PyCharm."
    )

    add_heading(doc, "Option A — Clone inside PyCharm (Easiest)", level=2)

    steps_a = [
        ("Open PyCharm",
         ["Launch PyCharm from your applications menu or desktop shortcut.",
          "On the Welcome screen you will see a button labelled 'Get from VCS' — click it."]),
        ("Enter the Repository URL",
         ["In the 'URL' field type or paste:",
          "    https://github.com/emorrow-hcg/event-request.git",
          "In the 'Directory' field, choose where on your computer to save the project.",
          "    Example: C:\\Projects\\event-request  (Windows)",
          "             /Users/yourname/Projects/event-request  (Mac/Linux)"]),
        ("Click Clone",
         ["PyCharm will download the project. This takes about 30 seconds.",
          "When prompted 'Open the project?' — click Yes."]),
    ]
    for n, (title, details) in enumerate(steps_a, 1):
        add_step(doc, n, title, details)

    add_heading(doc, "Option B — Clone using the Terminal", level=2)
    add_body(doc, "If you prefer using the command line:")
    add_code(doc, "git clone https://github.com/emorrow-hcg/event-request.git")
    add_body(doc, "Then open PyCharm and go to  File → Open  and select the event-request folder.")
    add_divider(doc)

    # ── SECTION 4: Python Interpreter ───────────────────────────────────────
    add_heading(doc, "4. Setting Up the Python Interpreter", level=1)
    add_body(doc,
        "PyCharm needs to know which version of Python to use. We will create a "
        "virtual environment — a self-contained copy of Python just for this project. "
        "This keeps the project's packages separate from other Python projects."
    )

    steps_py = [
        ("Open Settings",
         ["On Windows/Linux: go to  File → Settings",
          "On Mac: go to  PyCharm → Settings  (or press Cmd + , )"]),
        ("Navigate to the Interpreter",
         ["In the left panel, click  Project: event-request",
          "Then click  Python Interpreter"]),
        ("Add a New Interpreter",
         ["Click the gear icon (⚙) on the right side",
          "Select  Add Interpreter → Add Local Interpreter"]),
        ("Create a Virtualenv",
         ["Select  Virtualenv Environment  on the left",
          "Make sure  New environment  is selected",
          "In the  Base interpreter  dropdown, select  Python 3.13",
          "    If you don't see Python 3.13, click the three dots (...) and browse to",
          "    where you installed Python (e.g. C:\\Python313\\python.exe on Windows)",
          "Leave the location as the default  (it will be inside your project folder)",
          "Click  OK"]),
        ("Apply the Settings",
         ["PyCharm will create the virtual environment. This takes about 10 seconds.",
          "Click  OK  to close Settings."]),
    ]
    for n, (title, details) in enumerate(steps_py, 1):
        add_step(doc, n, title, details)

    add_tip(doc,
        "A virtual environment (venv) is like a clean sandbox. It means packages installed "
        "for this project won't interfere with other Python projects on your computer.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 5: Installing Dependencies ───────────────────────────────────
    add_heading(doc, "5. Installing the Required Packages", level=1)
    add_body(doc,
        "The project depends on several Python packages (FastAPI, Anthropic, etc.). "
        "PyCharm may detect these automatically and offer to install them. "
        "Follow the steps below to install them manually if needed."
    )

    steps_dep = [
        ("Open the PyCharm Terminal",
         ["At the bottom of the PyCharm window, click the  Terminal  tab.",
          "You should see a prompt that looks like:  (venv) C:\\...\\event-request>",
          "The  (venv)  prefix confirms your virtual environment is active."]),
        ("Run the Install Command",
         ["Type the following command exactly and press Enter:"]),
    ]
    for n, (title, details) in enumerate(steps_dep, 1):
        add_step(doc, n, title, details)

    add_code(doc, "pip install -r backend/requirements.txt")

    add_body(doc,
        "You will see a list of packages being downloaded. Wait until you see "
        "'Successfully installed...' before moving on."
    )
    add_tip(doc,
        "If PyCharm shows a yellow banner at the top saying 'Package requirements are not satisfied', "
        "you can also click the  Install requirements  button in that banner instead of using the terminal.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 6: API Key ───────────────────────────────────────────────────
    add_heading(doc, "6. Adding Your Anthropic API Key", level=1)
    add_body(doc,
        "The app uses Claude AI to optimize resumes. You need an API key so the app "
        "can communicate with the AI. Never share this key with anyone."
    )

    steps_key = [
        ("Open Run Configurations",
         ["At the top of PyCharm, look for the green  ▶  (Play) button",
          "Click the dropdown arrow  ▼  next to it",
          "Select  Edit Configurations..."]),
        ("Find the FastAPI Server Configuration",
         ["In the left panel of the dialog that appears, click  Run FastAPI Server",
          "(This configuration was pre-created for you — it's ready to use)"]),
        ("Add the Environment Variable",
         ["Find the  Environment variables  field",
          "Click the folder icon on the right side of that field",
          "Click the  +  button to add a new variable",
          "In the  Name  column type:   ANTHROPIC_API_KEY",
          "In the  Value  column paste your API key  (starts with sk-ant-...)",
          "Click  OK  to close the environment variable window",
          "Click  OK  to close Run Configurations"]),
    ]
    for n, (title, details) in enumerate(steps_key, 1):
        add_step(doc, n, title, details)

    add_tip(doc,
        "WARNING: Never paste your API key directly into code files or commit it to GitHub. "
        "Always use environment variables. Your key gives access to paid AI services.",
        kind='WARNING'
    )
    add_divider(doc)

    # ── SECTION 7: Running the App ───────────────────────────────────────────
    add_heading(doc, "7. Running the Application", level=1)
    add_body(doc, "Now you are ready to start the app.")

    steps_run = [
        ("Select the Run Configuration",
         ["Click the dropdown arrow  ▼  next to the green ▶ button at the top",
          "Select  Run FastAPI Server  from the list"]),
        ("Start the Server",
         ["Click the green  ▶  (Play) button",
          "PyCharm will open a Run panel at the bottom of the screen",
          "Wait until you see a line that says:",
          "    Uvicorn running on http://0.0.0.0:8000 (Press CTRL+C to quit)"]),
        ("Open the App in Your Browser",
         ["Open any web browser (Chrome, Edge, Firefox, Safari)",
          "Type the following in the address bar and press Enter:",
          "    http://localhost:8000",
          "You should see the ATS Resume Optimizer app."]),
    ]
    for n, (title, details) in enumerate(steps_run, 1):
        add_step(doc, n, title, details)

    add_tip(doc,
        "To stop the server, click the red  ■  (Stop) button in the Run panel, "
        "or press  Ctrl + F2  on Windows/Linux or  Cmd + F2  on Mac.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 8: Using the App ─────────────────────────────────────────────
    add_heading(doc, "8. Using the App — Step by Step", level=1)

    add_heading(doc, "Step 1 of 3 — Load a Job Description", level=2)
    add_body(doc, "Choose one of three ways to provide a job description:")
    add_bullet(doc, "URL Tab: Paste a link to a job posting (e.g. from LinkedIn or Indeed) and click Fetch & Analyze")
    add_bullet(doc, "Upload File Tab: Click 'Drop PDF/DOCX/TXT here' and select a saved job description file")
    add_bullet(doc, "Paste Text Tab: Copy and paste the job description text, then click Analyze")
    add_body(doc, "After a few seconds, the Extracted Requirements section will appear showing skills, keywords, and qualifications pulled from the job.")

    add_heading(doc, "Step 2 of 3 — Provide Your Resume", level=2)
    add_body(doc, "Scroll down to the 'Your Resume / Experience' section. Choose one option:")
    add_bullet(doc, "Upload Resume: Upload your existing resume as a PDF, DOCX, or TXT file")
    add_bullet(doc, "Paste Resume: Copy and paste your resume text, then click Optimize Resume")
    add_bullet(doc, "Build from Scratch: Fill in the form with your name, work history, and skills, then click Build & Optimize Resume")

    add_heading(doc, "Step 3 of 3 — Review the Results", level=2)
    add_body(doc, "After 20-60 seconds the Results section appears:")
    add_bullet(doc, "Before / After Score: See how much your ATS score improved (0-100 scale)")
    add_bullet(doc, "Score Breakdown: Keyword, Skills, Experience, and Education scores")
    add_bullet(doc, "Optimized Resume tab: Your rewritten resume — click Copy to Clipboard or Download .txt")
    add_bullet(doc, "Gap Analysis tab: See exactly which keywords and skills were missing and added")

    add_tip(doc,
        "A score of 75 or above is considered good for most ATS systems. "
        "The AI only uses experience you provided — it never invents qualifications.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 9: Running Tests ─────────────────────────────────────────────
    add_heading(doc, "9. Running the Tests", level=1)
    add_body(doc,
        "The project includes automated tests to verify the code works correctly. "
        "You can run them inside PyCharm."
    )

    steps_test = [
        ("Use the Pre-Built Test Configuration",
         ["Click the dropdown arrow  ▼  next to the green ▶ button",
          "Select  pytest",
          "Click the green  ▶  button",
          "PyCharm will open a test results panel showing all 9 tests"]),
        ("Or Right-Click to Run Individual Tests",
         ["In the Project panel on the left, expand  backend → tests",
          "Right-click on  test_api.py  or  test_parser.py",
          "Select  Run 'pytest in test_api.py'"]),
    ]
    for n, (title, details) in enumerate(steps_test, 1):
        add_step(doc, n, title, details)

    add_tip(doc,
        "All 9 tests should show green checkmarks. If a test fails, read the error message "
        "in the bottom panel — it will tell you exactly what went wrong.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 10: Troubleshooting ──────────────────────────────────────────
    add_heading(doc, "10. Troubleshooting", level=1)

    problems = [
        ("I see 'ANTHROPIC_API_KEY not configured'",
         "You need to add your API key to the Run Configuration (see Section 6)."),
        ("The browser shows 'This site can't be reached'",
         "Make sure the server is running — you should see the green ▶ active in PyCharm. "
         "Try http://127.0.0.1:8000 instead of localhost."),
        ("pip install fails with 'command not found'",
         "Make sure you are in the Terminal tab inside PyCharm (not an external terminal) "
         "and that (venv) appears in the prompt."),
        ("I don't see Python 3.13 in the interpreter list",
         "Go back to python.org and install Python 3.13. During installation on Windows, "
         "tick the 'Add Python to PATH' checkbox."),
        ("The optimization takes too long / times out",
         "This is normal during high AI traffic. Wait up to 90 seconds. "
         "If it fails, try again — it's usually a temporary issue."),
        ("PDF returns empty or garbled text",
         "The PDF may be a scanned image (not real text). Open it, select all text, copy it, "
         "and use the 'Paste Text' option instead."),
    ]

    for problem, solution in problems:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(6)
        qr = p.add_run(f"Q: {problem}")
        qr.bold = True
        qr.font.size = Pt(10.5)
        sp = doc.add_paragraph()
        sp.paragraph_format.left_indent = Inches(0.5)
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(8)
        sr = sp.add_run(f"A: {solution}")
        sr.font.size = Pt(10.5)

    add_divider(doc)

    # ── SECTION 11: Quick Reference ──────────────────────────────────────────
    add_heading(doc, "11. Quick Reference", level=1)

    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = 'Table Grid'
    hdr = tbl.rows[0].cells
    set_cell_bg(hdr[0], MID_BLUE)
    set_cell_bg(hdr[1], MID_BLUE)
    for cell, text in zip(hdr, ['Action', 'How']):
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(10)

    rows = [
        ("Start the server", "Select 'Run FastAPI Server' → click ▶"),
        ("Stop the server", "Click ■ in the Run panel, or press Ctrl+F2"),
        ("Open the app", "Go to http://localhost:8000 in your browser"),
        ("Run all tests", "Select 'pytest' configuration → click ▶"),
        ("Open terminal", "Click the Terminal tab at the bottom of PyCharm"),
        ("Install packages", "Terminal: pip install -r backend/requirements.txt"),
        ("Pull latest code", "Terminal: git pull origin main"),
        ("View API docs", "Go to http://localhost:8000/docs in your browser"),
    ]
    for action, how in rows:
        row = tbl.add_row().cells
        row[0].paragraphs[0].add_run(action).font.size = Pt(10)
        row[1].paragraphs[0].add_run(how).font.size = Pt(10)

    doc.add_paragraph()
    final = doc.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = final.add_run("Questions? Open an issue at github.com/emorrow-hcg/event-request/issues")
    fr.font.size = Pt(10)
    fr.font.color.rgb = ACCENT

    return doc


# ═══════════════════════════════════════════════════════════════════════════════
# VS CODE SOP
# ═══════════════════════════════════════════════════════════════════════════════

def build_vscode_sop():
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    title_page(
        doc,
        "ATS Resume Optimizer",
        "Step-by-Step Setup & Usage Guide",
        "Visual Studio Code Edition"
    )

    # ── SECTION 1: Overview ──────────────────────────────────────────────────
    add_heading(doc, "1. Overview", level=1)
    add_body(doc,
        "This guide walks you through every step needed to get the ATS Resume Optimizer "
        "running on your computer using Visual Studio Code (VS Code) — a free, lightweight, "
        "and very popular code editor. No prior experience is required. "
        "Follow each step in order and you will have the application running in under 20 minutes."
    )
    add_tip(doc,
        "ATS stands for Applicant Tracking System. This tool reads a job description, "
        "scores your resume against it, and uses AI to rewrite your resume so it scores higher.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 2: Prerequisites ─────────────────────────────────────────────
    add_heading(doc, "2. What You Need Before You Start", level=1)
    add_body(doc, "Install the following software on your computer before opening VS Code:")

    prereqs = [
        ("Python 3.13 (latest)",
         "https://www.python.org/downloads/",
         "Download and run the installer. On Windows, tick 'Add Python to PATH' before clicking Install. "
         "On Mac you may also use: brew install python@3.13"),
        ("Git",
         "https://git-scm.com/downloads",
         "Needed to download (clone) the project code from GitHub."),
        ("Visual Studio Code (free)",
         "https://code.visualstudio.com/",
         "Download the installer for your operating system and run it."),
        ("VS Code Python Extension",
         "Search 'Python' in VS Code Extensions (see Section 3)",
         "Published by Microsoft. Adds Python language support, linting, and the debugger."),
        ("An Anthropic API Key",
         "https://console.anthropic.com",
         "Sign in, go to API Keys, click Create Key, and copy the key somewhere safe."),
    ]

    for i, (name, url, note) in enumerate(prereqs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        nr = p.add_run(f"  {i}.  ")
        nr.bold = True
        nr.font.color.rgb = ACCENT
        nr.font.size = Pt(11)
        br = p.add_run(name)
        br.bold = True
        br.font.size = Pt(11)
        urlp = doc.add_paragraph()
        urlp.paragraph_format.left_indent = Inches(0.6)
        urlp.paragraph_format.space_before = Pt(0)
        urlp.paragraph_format.space_after = Pt(2)
        ur = urlp.add_run(f"Download: {url}")
        ur.font.size = Pt(10)
        ur.font.color.rgb = ACCENT
        np_ = doc.add_paragraph()
        np_.paragraph_format.left_indent = Inches(0.6)
        np_.paragraph_format.space_before = Pt(0)
        np_.paragraph_format.space_after = Pt(8)
        nr2 = np_.add_run(note)
        nr2.font.size = Pt(10)

    add_tip(doc,
        "To verify Python is installed: open a terminal and type  python --version  "
        "You should see 'Python 3.13.x'. On Mac/Linux you may need  python3 --version",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 3: Install Extensions ───────────────────────────────────────
    add_heading(doc, "3. Installing VS Code Extensions", level=1)
    add_body(doc,
        "Extensions add extra features to VS Code. You need to install two extensions "
        "before working with this project."
    )

    exts = [
        ("Python (by Microsoft)",
         "Click the Extensions icon on the left sidebar (it looks like four squares).\n"
         "Search for: Python\n"
         "Click on the result published by Microsoft.\n"
         "Click Install."),
        ("Python Debugger (by Microsoft)",
         "In the same Extensions panel, search for: Python Debugger\n"
         "Install the one published by Microsoft."),
        ("REST Client (optional but recommended)",
         "Search for: REST Client (by Huachao Mao)\n"
         "This lets you test the API endpoints directly inside VS Code using the api.http file."),
    ]
    for i, (name, steps) in enumerate(exts, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(6)
        nr = p.add_run(f"  Extension {i}: ")
        nr.bold = True
        nr.font.color.rgb = ACCENT
        nr.font.size = Pt(11)
        br = p.add_run(name)
        br.bold = True
        br.font.size = Pt(11)
        for line in steps.split('\n'):
            lp = doc.add_paragraph(line)
            lp.paragraph_format.left_indent = Inches(0.6)
            lp.paragraph_format.space_before = Pt(1)
            lp.paragraph_format.space_after = Pt(1)
            for r in lp.runs:
                r.font.size = Pt(10.5)

    add_divider(doc)

    # ── SECTION 4: Getting the Code ──────────────────────────────────────────
    add_heading(doc, "4. Getting the Project Code", level=1)
    add_body(doc, "You need to copy (clone) the project from GitHub to your computer.")

    add_heading(doc, "Option A — Clone using VS Code (Easiest)", level=2)

    steps_a = [
        ("Open the Command Palette",
         ["Press  Ctrl + Shift + P  (Windows/Linux) or  Cmd + Shift + P  (Mac)",
          "This opens the Command Palette — a search box for VS Code commands"]),
        ("Run the Clone Command",
         ["Type:  Git: Clone  and press Enter",
          "Paste the repository URL:",
          "    https://github.com/emorrow-hcg/event-request.git",
          "Press Enter"]),
        ("Choose Where to Save the Project",
         ["A file browser will appear — navigate to where you want to save the project",
          "    Example: C:\\Projects  (Windows) or  /Users/yourname/Projects  (Mac)",
          "Click  Select Repository Location"]),
        ("Open the Cloned Project",
         ["VS Code will ask 'Would you like to open the cloned repository?'",
          "Click  Open"]),
    ]
    for n, (title, details) in enumerate(steps_a, 1):
        add_step(doc, n, title, details)

    add_heading(doc, "Option B — Clone using the Terminal", level=2)
    add_body(doc, "Open any terminal on your computer and run:")
    add_code(doc, "git clone https://github.com/emorrow-hcg/event-request.git")
    add_body(doc, "Then in VS Code go to  File → Open Folder  and select the event-request folder.")
    add_divider(doc)

    # ── SECTION 5: Python Interpreter ────────────────────────────────────────
    add_heading(doc, "5. Setting Up the Python Environment", level=1)
    add_body(doc,
        "VS Code needs to know which Python to use. We will create a virtual environment "
        "— a self-contained Python sandbox just for this project."
    )

    steps_py = [
        ("Open the VS Code Terminal",
         ["Go to  Terminal → New Terminal  from the menu bar",
          "A terminal panel will open at the bottom of VS Code",
          "Make sure the path shown ends with  event-request"]),
        ("Create a Virtual Environment",
         ["Type the following command and press Enter:"]),
    ]
    for n, (title, details) in enumerate(steps_py, 1):
        add_step(doc, n, title, details)

    add_body(doc, "Windows:")
    add_code(doc, "py -3.13 -m venv .venv")
    add_body(doc, "Mac / Linux:")
    add_code(doc, "python3.13 -m venv .venv")

    steps_py2 = [
        ("Select the Python Interpreter",
         ["Press  Ctrl + Shift + P  (Cmd + Shift + P on Mac)",
          "Type:  Python: Select Interpreter  and press Enter",
          "A list of Python versions appears — select the one that shows  .venv  in its path",
          "    It will look like:  Python 3.13.x ('.venv': venv)  ./venv/bin/python",
          "If you don't see it, click  Enter interpreter path  and browse to:",
          "    Windows:  .venv\\Scripts\\python.exe",
          "    Mac/Linux:  .venv/bin/python"]),
        ("Activate the Virtual Environment in the Terminal",
         ["VS Code may do this automatically. To activate manually:"]),
    ]
    for n, (title, details) in enumerate(steps_py2, 3):
        add_step(doc, n, title, details)

    add_body(doc, "Windows:")
    add_code(doc, ".venv\\Scripts\\activate")
    add_body(doc, "Mac / Linux:")
    add_code(doc, "source .venv/bin/activate")
    add_body(doc, "After activation, the terminal prompt will start with  (.venv)")

    add_tip(doc,
        "A virtual environment (venv) is a safe sandbox. Packages installed here won't "
        "affect other Python projects on your computer — and vice versa.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 6: Install Dependencies ──────────────────────────────────────
    add_heading(doc, "6. Installing the Required Packages", level=1)
    add_body(doc,
        "In the VS Code terminal (with (.venv) shown in the prompt), run:"
    )
    add_code(doc, "pip install -r backend/requirements.txt")
    add_body(doc,
        "Wait until you see 'Successfully installed...' before moving on. "
        "This downloads all the libraries the app needs (FastAPI, Anthropic, etc.)."
    )
    add_tip(doc,
        "If pip says 'not found', make sure (.venv) is in the terminal prompt. "
        "If not, re-run the activate command from Section 5 Step 4.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 7: API Key ────────────────────────────────────────────────────
    add_heading(doc, "7. Adding Your Anthropic API Key", level=1)
    add_body(doc,
        "VS Code reads environment variables from a file called  .env  in the project folder. "
        "Follow these steps to create it."
    )

    steps_key = [
        ("Create the .env File",
         ["In VS Code's Explorer panel (left sidebar), you will see the project files",
          "Look for a file called  .env.example  — right-click it",
          "Select  Copy  then right-click the project folder and select  Paste",
          "Rename the copy to:  .env  (remove '.example' from the name)"]),
        ("Edit the .env File",
         ["Double-click  .env  to open it",
          "You will see:   ANTHROPIC_API_KEY=your_anthropic_api_key_here",
          "Replace  your_anthropic_api_key_here  with your actual key (starts with sk-ant-...)",
          "Save the file:  Ctrl + S  (Cmd + S on Mac)"]),
        ("Verify the Launch Configuration Reads It",
         ["The pre-built launch configuration in this project is set to read  .env  automatically",
          "You don't need to do anything else — the key will be picked up when you run the app"]),
    ]
    for n, (title, details) in enumerate(steps_key, 1):
        add_step(doc, n, title, details)

    add_tip(doc,
        "WARNING: The .env file contains your secret API key. "
        "Never commit it to GitHub. It is already listed in .gitignore so Git will ignore it automatically.",
        kind='WARNING'
    )
    add_divider(doc)

    # ── SECTION 8: VS Code Launch Config ─────────────────────────────────────
    add_heading(doc, "8. Creating the VS Code Launch Configuration", level=1)
    add_body(doc,
        "VS Code uses a  launch.json  file to know how to run the project. "
        "Follow these steps to create it."
    )

    steps_launch = [
        ("Open the Run and Debug Panel",
         ["Click the bug + play icon on the left sidebar  (or press  Ctrl + Shift + D)"]),
        ("Create launch.json",
         ["Click  'create a launch.json file'  link",
          "VS Code asks to select a debugger — choose  Python Debugger",
          "Then choose  FastAPI"]),
        ("Edit the Configuration",
         ["VS Code opens  .vscode/launch.json",
          "Replace the entire contents with the following:"]),
    ]
    for n, (title, details) in enumerate(steps_launch, 1):
        add_step(doc, n, title, details)

    add_code(doc, '''{
  "version": "0.2.0",
  "configurations": [
    {
      "name": "Run FastAPI Server",
      "type": "debugpy",
      "request": "launch",
      "module": "uvicorn",
      "args": ["main:app", "--reload", "--host", "0.0.0.0", "--port", "8000"],
      "cwd": "${workspaceFolder}/backend",
      "envFile": "${workspaceFolder}/.env",
      "console": "integratedTerminal"
    },
    {
      "name": "pytest",
      "type": "debugpy",
      "request": "launch",
      "module": "pytest",
      "args": ["tests/", "-v"],
      "cwd": "${workspaceFolder}/backend",
      "envFile": "${workspaceFolder}/.env",
      "console": "integratedTerminal"
    }
  ]
}''')

    add_step(doc, 4, "Save the File",
             ["Press  Ctrl + S  (Cmd + S on Mac) to save  launch.json"])

    add_tip(doc,
        "You only need to create launch.json once. From now on, the Run and Debug panel "
        "will always show your configurations.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 9: Running the App ────────────────────────────────────────────
    add_heading(doc, "9. Running the Application", level=1)

    steps_run = [
        ("Open the Run and Debug Panel",
         ["Click the bug + play icon on the left sidebar  (or press  Ctrl + Shift + D)"]),
        ("Select the Configuration",
         ["At the top of the panel, click the dropdown and select  Run FastAPI Server"]),
        ("Start the Server",
         ["Click the green  ▶  (Play) button next to the dropdown",
          "The terminal at the bottom will show output. Wait until you see:",
          "    Uvicorn running on http://0.0.0.0:8000 (Press CTRL+C to quit)"]),
        ("Open the App in Your Browser",
         ["Open any web browser",
          "Type  http://localhost:8000  in the address bar and press Enter",
          "You should see the ATS Resume Optimizer app"]),
    ]
    for n, (title, details) in enumerate(steps_run, 1):
        add_step(doc, n, title, details)

    add_tip(doc,
        "To stop the server, press  Shift + F5  in VS Code, "
        "or click the red  ■  (Stop) button in the debug toolbar at the top.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 10: Using the App ─────────────────────────────────────────────
    add_heading(doc, "10. Using the App — Step by Step", level=1)

    add_heading(doc, "Step 1 of 3 — Load a Job Description", level=2)
    add_body(doc, "Choose one of three ways to provide a job description:")
    add_bullet(doc, "URL Tab: Paste a link to a job posting (e.g. from LinkedIn or Indeed) and click Fetch & Analyze")
    add_bullet(doc, "Upload File Tab: Click 'Drop PDF/DOCX/TXT here' and select a saved job description file")
    add_bullet(doc, "Paste Text Tab: Copy and paste the job description text, then click Analyze")
    add_body(doc, "The Extracted Requirements section will appear showing skills, keywords, and qualifications from the job.")

    add_heading(doc, "Step 2 of 3 — Provide Your Resume", level=2)
    add_body(doc, "Scroll down to the 'Your Resume / Experience' section. Choose one option:")
    add_bullet(doc, "Upload Resume: Upload your existing resume as a PDF, DOCX, or TXT file")
    add_bullet(doc, "Paste Resume: Copy and paste your resume text, then click Optimize Resume")
    add_bullet(doc, "Build from Scratch: Fill in the guided form with your details, then click Build & Optimize Resume")

    add_heading(doc, "Step 3 of 3 — Review the Results", level=2)
    add_body(doc, "After 20-60 seconds the Results section appears:")
    add_bullet(doc, "Before / After Score: See how much your ATS score improved (0-100 scale)")
    add_bullet(doc, "Score Breakdown: Keyword, Skills, Experience, and Education scores with progress bars")
    add_bullet(doc, "Optimized Resume tab: Your rewritten resume — click Copy to Clipboard or Download .txt")
    add_bullet(doc, "Gap Analysis tab: See which keywords and skills were missing and which were matched")

    add_tip(doc,
        "A score of 75 or above is considered good for most ATS systems. "
        "The AI only uses experience you provided — it never invents qualifications.",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 11: Running Tests ─────────────────────────────────────────────
    add_heading(doc, "11. Running the Tests", level=1)
    add_body(doc,
        "The project includes 9 automated tests. Running them confirms everything is working."
    )

    steps_test = [
        ("Open Run and Debug",
         ["Click the bug + play icon on the left sidebar  (Ctrl + Shift + D)"]),
        ("Select the pytest Configuration",
         ["In the dropdown at the top, select  pytest",
          "Click the green  ▶  button"]),
        ("View Results",
         ["The terminal shows each test running",
          "All 9 tests should show green  PASSED  next to them",
          "If any test fails, read the error message — it will tell you what to fix"]),
    ]
    for n, (title, details) in enumerate(steps_test, 1):
        add_step(doc, n, title, details)

    add_tip(doc,
        "You can also run tests from the terminal:  python -m pytest backend/tests/ -v",
        kind='TIP'
    )
    add_divider(doc)

    # ── SECTION 12: Troubleshooting ───────────────────────────────────────────
    add_heading(doc, "12. Troubleshooting", level=1)

    problems = [
        ("I see 'ANTHROPIC_API_KEY not configured'",
         "Your .env file is missing or the key value is still the placeholder text. "
         "Open .env and replace 'your_anthropic_api_key_here' with your actual key from console.anthropic.com."),
        ("The browser shows 'This site can't be reached'",
         "Make sure the server is running — the terminal should show 'Uvicorn running on http://0.0.0.0:8000'. "
         "Try http://127.0.0.1:8000 instead."),
        ("'python3.13' is not recognized / command not found",
         "Python 3.13 is not installed or not in your PATH. "
         "Re-install from python.org — on Windows, tick 'Add Python to PATH' during installation."),
        ("(.venv) is not shown in the terminal prompt",
         "The virtual environment is not activated. Run the activate command from Section 5 Step 4. "
         "On Windows: .venv\\Scripts\\activate   On Mac/Linux: source .venv/bin/activate"),
        ("VS Code doesn't find the .venv interpreter",
         "Press Ctrl+Shift+P → 'Python: Select Interpreter' → 'Enter interpreter path' "
         "→ browse to .venv/Scripts/python.exe (Windows) or .venv/bin/python (Mac/Linux)."),
        ("The optimization takes too long",
         "This is normal during high AI traffic periods. Wait up to 90 seconds. "
         "If it times out, just try again."),
        ("PDF returns empty text",
         "The PDF may be image-based (scanned). Open it, select all text, copy, and use the Paste Text option."),
    ]

    for problem, solution in problems:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(6)
        qr = p.add_run(f"Q: {problem}")
        qr.bold = True
        qr.font.size = Pt(10.5)
        sp = doc.add_paragraph()
        sp.paragraph_format.left_indent = Inches(0.5)
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(8)
        sr = sp.add_run(f"A: {solution}")
        sr.font.size = Pt(10.5)

    add_divider(doc)

    # ── SECTION 13: Quick Reference ───────────────────────────────────────────
    add_heading(doc, "13. Quick Reference", level=1)

    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = 'Table Grid'
    hdr = tbl.rows[0].cells
    set_cell_bg(hdr[0], MID_BLUE)
    set_cell_bg(hdr[1], MID_BLUE)
    for cell, text in zip(hdr, ['Action', 'How']):
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(10)

    rows = [
        ("Start the server",         "Run & Debug panel → 'Run FastAPI Server' → ▶"),
        ("Stop the server",           "Press Shift + F5  or click ■ in the debug toolbar"),
        ("Open the app",              "Go to http://localhost:8000 in your browser"),
        ("Run all tests",             "Run & Debug panel → 'pytest' → ▶"),
        ("Open terminal",             "Terminal → New Terminal  or  Ctrl + ` (backtick)"),
        ("Activate virtual env",      "Windows: .venv\\Scripts\\activate   Mac/Linux: source .venv/bin/activate"),
        ("Install packages",          "pip install -r backend/requirements.txt"),
        ("Pull latest code",          "git pull origin main  (in terminal)"),
        ("View API docs",             "Go to http://localhost:8000/docs in your browser"),
        ("Open Command Palette",      "Ctrl + Shift + P  (Cmd + Shift + P on Mac)"),
    ]
    for action, how in rows:
        row = tbl.add_row().cells
        row[0].paragraphs[0].add_run(action).font.size = Pt(10)
        row[1].paragraphs[0].add_run(how).font.size = Pt(10)

    doc.add_paragraph()
    final = doc.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = final.add_run("Questions? Open an issue at github.com/emorrow-hcg/event-request/issues")
    fr.font.size = Pt(10)
    fr.font.color.rgb = ACCENT

    return doc


# ── Generate both files ───────────────────────────────────────────────────────
if __name__ == "__main__":
    import os
    out_dir = os.path.join(os.path.dirname(__file__), "..", "docs")
    os.makedirs(out_dir, exist_ok=True)

    pycharm_path = os.path.join(out_dir, "SOP_PyCharm_Setup_Guide.docx")
    vscode_path  = os.path.join(out_dir, "SOP_VSCode_Setup_Guide.docx")

    print("Building PyCharm SOP...")
    build_pycharm_sop().save(pycharm_path)
    print(f"  Saved → {pycharm_path}")

    print("Building VS Code SOP...")
    build_vscode_sop().save(vscode_path)
    print(f"  Saved → {vscode_path}")

    print("Done.")
